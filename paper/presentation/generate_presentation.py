"""Generate an editable 25-minute academic presentation for the project."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = Path(__file__).resolve().parent / "ietf_ambiguity_detection_presentation.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
NAVY = RGBColor(24, 52, 82)
BLUE = RGBColor(45, 101, 150)
TEAL = RGBColor(28, 133, 133)
ORANGE = RGBColor(218, 139, 53)
RED = RGBColor(166, 65, 62)
INK = RGBColor(38, 49, 59)
MUTED = RGBColor(91, 105, 116)
PALE = RGBColor(242, 247, 250)
PALE_BLUE = RGBColor(231, 240, 247)
PALE_TEAL = RGBColor(231, 245, 243)
PALE_ORANGE = RGBColor(251, 242, 227)
PALE_RED = RGBColor(250, 237, 235)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(214, 224, 231)
FONT = "Aptos"


def inches(value: float) -> int:
    return Inches(value)


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
):
    shape = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = inches(margin)
    frame.margin_right = inches(margin)
    frame.margin_top = inches(margin)
    frame.margin_bottom = inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(5)
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
    return shape


def add_box(slide, x, y, w, h, fill=WHITE, outline=LINE, radius=True, width=1.1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, inches(x), inches(y), inches(w), inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = outline
    shape.line.width = Pt(width)
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, inches(x1), inches(y1), inches(x2), inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def add_circle(slide, x, y, d, fill, text, size=18, color=WHITE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, inches(x), inches(y), inches(d), inches(d)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    add_text(
        slide, x, y + 0.005, d, d - 0.01, text, size=size, color=color, bold=True,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0,
    )
    return shape


def add_header(slide, title, section=None):
    add_text(slide, 0.62, 0.34, 11.9, 0.48, title, size=27, color=NAVY, bold=True)
    add_line(slide, 0.66, 0.96, 12.66, 0.96, color=BLUE, width=1.6)
    if section:
        add_text(slide, 10.6, 0.1, 2.0, 0.22, section.upper(), size=9, color=TEAL,
                 bold=True, align=PP_ALIGN.RIGHT)


def add_footer(slide, page, source=""):
    if source:
        add_text(slide, 0.66, 7.14, 10.8, 0.18, source, size=8.5, color=MUTED)
    add_text(slide, 11.75, 7.1, 0.85, 0.22, f"{page:02d}  |  18", size=9,
             color=MUTED, align=PP_ALIGN.RIGHT)


def new_slide(prs, title=None, section=None, source="", page=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    if title:
        add_header(slide, title, section)
    if page is not None:
        add_footer(slide, page, source)
    return slide


def add_card(slide, x, y, w, h, title, body, fill=PALE, accent=BLUE,
             title_size=17, body_size=14):
    add_box(slide, x, y, w, h, fill=fill, outline=LINE)
    add_box(slide, x, y, 0.09, h, fill=accent, outline=accent, radius=False, width=0)
    add_text(slide, x + 0.24, y + 0.2, w - 0.43, 0.38, title,
             size=title_size, color=NAVY, bold=True)
    add_text(slide, x + 0.24, y + 0.67, w - 0.45, h - 0.82, body,
             size=body_size, color=INK)


def add_metric(slide, x, y, w, h, value, label, color=BLUE):
    add_box(slide, x, y, w, h, fill=PALE_BLUE, outline=PALE_BLUE)
    add_text(slide, x + 0.1, y + 0.16, w - 0.2, 0.58, value, size=31,
             color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.12, y + 0.82, w - 0.24, h - 0.91, label, size=14,
             color=INK, align=PP_ALIGN.CENTER)


def add_bullet_list(slide, x, y, w, h, items, size=18, color=INK, gap=0.72):
    for index, item in enumerate(items):
        yy = y + index * gap
        add_circle(slide, x, yy + 0.04, 0.18, TEAL, "", size=6)
        add_text(slide, x + 0.32, yy, w - 0.32, gap - 0.06, item, size=size, color=color)


def add_bar_chart(slide, x, y, w, rows, max_value, baseline=None, row_h=0.62,
                  label_w=2.5, value_fmt="{:.3f}", highlight=None):
    chart_x = x + label_w
    chart_w = w - label_w - 0.65
    if baseline is not None:
        bx = chart_x + chart_w * baseline / max_value
        add_line(slide, bx, y - 0.04, bx, y + len(rows) * row_h - 0.08,
                 color=RED, width=1.2)
        add_text(slide, bx + 0.04, y - 0.27, 1.1, 0.22,
                 f"base {baseline:.3f}", size=9, color=RED)
    for i, (label, value, color) in enumerate(rows):
        yy = y + i * row_h
        add_text(slide, x, yy + 0.04, label_w - 0.12, 0.38, label,
                 size=14, color=INK, bold=(highlight == label))
        bar_w = max(0.06, chart_w * value / max_value)
        add_box(slide, chart_x, yy + 0.08, bar_w, 0.34, fill=color, outline=color,
                radius=False, width=0)
        add_text(slide, chart_x + bar_w + 0.09, yy + 0.05, 0.62, 0.35,
                 value_fmt.format(value), size=12, color=INK, bold=True)
    return chart_x, chart_w


def build_deck():
    prs = Presentation()
    prs.slide_width = inches(SLIDE_W)
    prs.slide_height = inches(SLIDE_H)
    prs.core_properties.title = "Leakage-Controlled Pre-Publication Ambiguity Detection"
    prs.core_properties.subject = "25-minute academic project presentation"
    prs.core_properties.author = "Queen's University School of Computing"
    prs.core_properties.keywords = "Internet-Drafts, ambiguity detection, ML, leakage control"

    # 1. Title
    slide = new_slide(prs)
    slide.background.fill.fore_color.rgb = PALE
    add_box(slide, 8.65, 0, 4.68, SLIDE_H, fill=NAVY, outline=NAVY, radius=False, width=0)
    for x, y, r, fill in [
        (9.1, 1.0, 0.18, TEAL), (10.6, 1.7, 0.24, ORANGE),
        (11.9, 0.92, 0.16, BLUE), (9.6, 3.1, 0.22, BLUE),
        (11.15, 3.8, 0.2, TEAL), (12.2, 5.5, 0.22, ORANGE),
        (9.35, 5.85, 0.14, TEAL),
    ]:
        add_circle(slide, x, y, r, fill, "", size=8)
    for a, b in [((9.28, 1.18), (10.7, 1.82)), ((10.82, 1.9), (11.98, 1.04)),
                 ((9.2, 1.2), (9.72, 3.2)), ((9.78, 3.2), (11.24, 3.9)),
                 ((11.35, 3.95), (12.28, 5.62)), ((9.75, 3.2), (9.42, 5.92))]:
        add_line(slide, *a, *b, color=RGBColor(91, 129, 158), width=1.6)
    add_text(slide, 0.85, 0.9, 7.2, 0.4, "RESEARCH PRESENTATION  |  25 MIN", size=12,
             color=TEAL, bold=True)
    add_text(slide, 0.82, 1.55, 7.3, 1.62,
             "Finding ambiguity\nbefore publication", size=35, color=NAVY, bold=True)
    add_text(slide, 0.88, 3.5, 7.0, 0.95,
             "Leakage-controlled, model-assisted prioritization of\nIETF Internet-Draft clauses",
             size=22, color=INK)
    add_text(slide, 0.9, 5.25, 7.15, 0.48,
             "Mostafa Elofy  ·  Mohamed Ghanem  ·  Mohamed Abdelkhalek  ·  Osama Elkhuribi",
             size=13, color=NAVY, bold=True)
    add_text(slide, 0.9, 5.8, 6.7, 0.42,
             "School of Computing, Queen's University", size=15, color=MUTED)
    add_text(slide, 0.9, 6.77, 6.7, 0.24,
             "Research-stage evidence — not a validated production detector", size=11,
             color=RED, bold=True)
    add_footer(slide, 1)

    # 2. Motivation
    slide = new_slide(prs, "One sentence can lead to different implementations",
                      "Motivation", page=2)
    add_text(slide, 0.9, 1.28, 11.5, 0.55,
             "Normative keywords are precise; the surrounding prose can still leave room for interpretation.",
             size=19, color=INK)
    add_box(slide, 3.65, 2.15, 6.0, 1.05, fill=PALE_BLUE, outline=BLUE, width=1.7)
    add_text(slide, 3.9, 2.34, 5.5, 0.62,
             "“A receiver SHOULD retry when appropriate.”", size=22, color=NAVY,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_line(slide, 6.65, 3.2, 6.65, 3.55, color=BLUE, width=1.5)
    add_line(slide, 2.75, 3.55, 10.55, 3.55, color=BLUE, width=1.5)
    for center, title, body, fill, accent in [
        (2.75, "Implementation A", "Retry immediately;\nstop after one failure.", PALE_TEAL, TEAL),
        (6.65, "Implementation B", "Retry with backoff;\ncontinue until a limit.", PALE_ORANGE, ORANGE),
        (10.55, "Implementation C", "Retry only for selected\nerror classes.", PALE_RED, RED),
    ]:
        add_line(slide, center, 3.55, center, 3.88, color=BLUE, width=1.3)
        add_box(slide, center - 1.58, 3.88, 3.16, 1.46, fill=fill, outline=accent)
        add_text(slide, center - 1.38, 4.08, 2.76, 0.32, title, size=16,
                 color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, center - 1.37, 4.55, 2.74, 0.58, body, size=15,
                 color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, 1.1, 5.85, 11.1, 0.57,
             "The goal is not to declare a clause defective. It is to help a human reviewer decide where to look first.",
             size=19, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 2, "Illustrative example; not a quoted clause from the corpus.")

    # 3. Research question
    slide = new_slide(prs, "The research question is about prioritization—not verdicts",
                      "Research question", page=3)
    add_box(slide, 0.95, 1.5, 11.4, 1.43, fill=PALE_BLUE, outline=PALE_BLUE)
    add_text(slide, 1.35, 1.78, 10.6, 0.9,
             "Can we rank draft clauses likely to be operationally ambiguous using only information available at that draft stage?",
             size=25, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE)
    add_card(slide, 1.0, 3.55, 3.55, 1.72, "Input",
             "A clause as it appeared in an Internet-Draft", fill=PALE, accent=BLUE)
    add_card(slide, 4.9, 3.55, 3.55, 1.72, "Output",
             "A ranked queue for human review", fill=PALE_TEAL, accent=TEAL)
    add_card(slide, 8.8, 3.55, 3.55, 1.72, "Boundary",
             "No automatic judgment of correctness", fill=PALE_ORANGE, accent=ORANGE)
    add_text(slide, 1.2, 5.85, 10.9, 0.62,
             "Success means useful review allocation under a defensible evaluation—not “AI understands standards.”",
             size=19, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 3, "Project framing: paper/main.tex, Introduction and Conclusion.")

    # 4. Challenge
    slide = new_slide(prs, "Three traps can make an ambiguity detector look better than it is",
                      "Problem framing", page=4)
    add_card(slide, 0.85, 1.55, 3.75, 3.28, "1  Future-information leakage",
             "Later revisions, final RFCs, and errata can explain what became problematic—but they were not available at draft time.",
             fill=PALE_RED, accent=RED, title_size=18, body_size=16)
    add_card(slide, 4.8, 1.55, 3.75, 3.28, "2  Repeated language",
             "Related clauses recur across a document's revision lineage. Random clause splits can place near-duplicates in train and test.",
             fill=PALE_ORANGE, accent=ORANGE, title_size=18, body_size=16)
    add_card(slide, 8.75, 1.55, 3.75, 3.28, "3  Rare, uncertain positives",
             "The latest DeBERTa cohort is 6.4% positive. Operational labels are AI-assisted and not expert-adjudicated gold.",
             fill=PALE_BLUE, accent=BLUE, title_size=18, body_size=16)
    add_text(slide, 1.0, 5.55, 11.3, 0.74,
             "So the central engineering problem is as much about information boundaries and evaluation design as model choice.",
             size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 4, "Source: paper/main.tex, Introduction; latest DeBERTa cohort.")

    # 5. Corpus/cohorts
    slide = new_slide(prs, "The project connects a corpus to several distinct studies",
                      "Data", page=5)
    add_metric(slide, 0.95, 1.35, 3.55, 1.4, "253", "Internet-Draft revisions")
    add_metric(slide, 4.9, 1.35, 3.55, 1.4, "12", "RFCs")
    add_metric(slide, 8.85, 1.35, 3.55, 1.4, "31", "verified technical errata")
    studies = [
        ("CLASSICAL ML", "2,406 clauses · 11 families", "158 positives · prevalence 6.57%", BLUE),
        ("QWEN / HYBRID", "2,142 clauses · 9 families", "126 positives · prevalence 5.88%", ORANGE),
        ("LATEST DEBERTA", "2,125 clauses · 9 families", "136 positives · prevalence 6.40%", TEAL),
    ]
    for i, (title, cohort, detail, accent) in enumerate(studies):
        y = 3.15 + i * 0.9
        add_box(slide, 1.2, y, 10.95, 0.72,
                fill=PALE if i != 1 else PALE_ORANGE, outline=LINE)
        add_text(slide, 1.45, y + 0.16, 2.45, 0.3, title, size=14,
                 color=accent, bold=True)
        add_text(slide, 4.0, y + 0.12, 3.6, 0.34, cohort, size=16,
                 color=NAVY, bold=True)
        add_text(slide, 7.8, y + 0.15, 3.95, 0.3, detail, size=14, color=INK)
    add_text(slide, 1.2, 6.14, 10.95, 0.42,
             "Different label snapshots and eligibility rules: these cohorts are not one shared leaderboard.",
             size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 5, "Source: README.md; paper/main.tex. Cohort metrics are study-specific.")

    # 6. Pipeline
    slide = new_slide(prs, "The pipeline makes the information boundary explicit",
                      "Method", page=6)
    steps = [
        ("1", "Public sources", "Drafts · RFCs\nerrata", BLUE),
        ("2", "Parse", "Documents →\nclauses", TEAL),
        ("3", "Track change", "Revision alignment\n& lineage", ORANGE),
        ("4", "Label", "AI-assisted\noperational reference", RED),
        ("5", "Model", "Rules · ML ·\nLLM · transformer", BLUE),
        ("6", "Evaluate", "Family-grouped\nout-of-fold", TEAL),
        ("7", "Prioritize", "Ranked human\nreview queue", ORANGE),
    ]
    lefts = [0.55, 2.38, 4.21, 6.04, 7.87, 9.70, 11.53]
    for (number, title, body, accent), x in zip(steps, lefts):
        add_box(slide, x, 2.12, 1.28, 2.1, fill=PALE, outline=LINE)
        add_circle(slide, x + 0.43, 1.73, 0.42, accent, number, size=15)
        add_text(slide, x + 0.08, 2.36, 1.12, 0.47, title, size=14,
                 color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.08, 3.0, 1.12, 0.83, body, size=12.5,
                 color=INK, align=PP_ALIGN.CENTER)
        if number != "7":
            add_text(slide, x + 1.34, 2.82, 0.4, 0.42, "›", size=28, color=BLUE,
                     bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 0.9, 4.72, 11.55, 0.82, fill=PALE_RED, outline=PALE_RED)
    add_text(slide, 1.16, 4.89, 11.05, 0.42,
             "Prediction-time boundary: no later draft, final RFC, erratum, label, vote, rationale, or split metadata as model input.",
             size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.2, 5.93, 10.9, 0.48,
             "Reserved TCP transfer validation and sealed DNS final testing remain unopened.",
             size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6, "Source: paper/figures/project_pipeline.png and paper/main.tex.")

    # 7. Labels
    slide = new_slide(prs, "The label is an operational reference—not a definition of truth",
                      "Measurement", page=7)
    add_box(slide, 0.95, 1.45, 5.4, 3.78, fill=PALE_BLUE, outline=PALE_BLUE)
    add_text(slide, 1.35, 1.78, 4.6, 0.48, "What the project has", size=21,
             color=NAVY, bold=True)
    add_bullet_list(slide, 1.35, 2.5, 4.55, 2.2, [
        "A repeatable annotation workflow",
        "Provenance and consensus rules",
        "AI-assisted reference labels",
    ], size=16, gap=0.72)
    add_box(slide, 6.95, 1.45, 5.4, 3.78, fill=PALE_ORANGE, outline=PALE_ORANGE)
    add_text(slide, 7.35, 1.78, 4.6, 0.48, "What remains missing", size=21,
             color=NAVY, bold=True)
    add_bullet_list(slide, 7.35, 2.5, 4.55, 2.25, [
        "Independent expert double annotation",
        "Adjudication against the full rubric",
        "A public, sealed gold test set",
    ], size=16, gap=0.72)
    add_text(slide, 1.25, 5.72, 10.9, 0.6,
             "Therefore, model scores measure agreement with this operational reference—not ambiguity accuracy.",
             size=19, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7, "Source: paper/main.tex, Construct Validity and Annotation sections.")

    # 8. Model families
    slide = new_slide(prs, "The project tests multiple model families—not one magic model",
                      "Models", page=8)
    add_card(slide, 0.9, 1.45, 5.45, 1.75, "Transparent rules",
             "Hand-set lexical and structural cues; an interpretable low-cost baseline.",
             fill=PALE_ORANGE, accent=ORANGE, title_size=17, body_size=15)
    add_card(slide, 6.95, 1.45, 5.45, 1.75, "Classical ML",
             "Logistic regression / SVM; engineered clause features + TF–IDF.",
             fill=PALE_BLUE, accent=BLUE, title_size=17, body_size=15)
    add_card(slide, 0.9, 3.55, 5.45, 1.75, "Transformers",
             "DeBERTa-v3-base; target-only fine-tuning and rank-based combinations.",
             fill=PALE_TEAL, accent=TEAL, title_size=17, body_size=15)
    add_card(slide, 6.95, 3.55, 5.45, 1.75, "Local LLM & hybrid",
             "Qwen prompting plus lexical / MiniLM semantic signals in separate studies.",
             fill=PALE, accent=NAVY, title_size=17, body_size=15)
    add_text(slide, 1.0, 5.9, 11.4, 0.52,
             "Revision-only signals and frozen embeddings were also examined; the tested versions were not consistently stronger.",
             size=17, color=INK, align=PP_ALIGN.CENTER)
    add_footer(slide, 8, "Source: paper/main.tex; results/final_summary/conclusions.md.")

    # 9. Evaluation
    slide = new_slide(prs, "Evaluation is grouped to test transfer across protocol families",
                      "Evaluation", page=9)
    add_text(slide, 0.95, 1.38, 5.2, 0.42, "WHAT THE DESIGN PROTECTS", size=13,
             color=TEAL, bold=True)
    add_bullet_list(slide, 1.0, 1.92, 5.2, 3.8, [
        "Keep related clauses from the same family together",
        "Fit preprocessing and feature selection inside training folds",
        "Use average precision (AP) under class imbalance",
        "Report fixed review-budget retrieval",
    ], size=16, gap=0.86)
    add_box(slide, 6.75, 1.5, 5.55, 3.75, fill=PALE, outline=LINE)
    add_text(slide, 7.12, 1.85, 4.85, 0.45, "Latest DeBERTa study", size=21,
             color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_metric(slide, 7.1, 2.57, 2.1, 1.55, "5-fold", "family-isolated OOF", color=TEAL)
    add_metric(slide, 9.82, 2.57, 2.1, 1.55, "1 seed", "run-to-run uncertainty remains", color=ORANGE)
    add_text(slide, 7.06, 4.53, 4.95, 0.4,
             "Development evaluation—not a sealed test.", size=16, color=RED,
             bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.12, 5.95, 11.0, 0.56,
             "No single score answers whether the system transfers to a new protocol family.",
             size=19, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 9, "Source: paper/main.tex, Evaluation Design and Results.")

    # 10. Classical result
    slide = new_slide(prs, "Classical text features give the strongest result in their own cohort",
                      "Results · classical ML", page=10)
    add_text(slide, 0.95, 1.25, 6.1, 0.52,
             "2,406 clauses  ·  158 positives  ·  11 families", size=18,
             color=MUTED, bold=True)
    rows = [
        ("Logistic: engineered + TF–IDF", 0.5225, BLUE),
        ("Linear SVM: engineered + TF–IDF", 0.5109, TEAL),
        ("Fixed engineered heuristic", 0.1049, ORANGE),
        ("Positive prevalence", 0.0657, MUTED),
    ]
    add_bar_chart(slide, 1.05, 2.15, 8.5, rows, max_value=0.58, baseline=0.0657,
                  row_h=0.73, label_w=3.7, value_fmt="{:.4f}", highlight=rows[0][0])
    add_box(slide, 9.8, 2.0, 2.55, 2.85, fill=PALE_BLUE, outline=PALE_BLUE)
    add_text(slide, 10.02, 2.3, 2.1, 0.5, "Key signal", size=17, color=TEAL,
             bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 10.0, 2.95, 2.15, 1.45,
             "Word-level patterns plus engineered features ranked positives above simple rules.",
             size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 1.1, 5.75, 11.1, 0.58,
             "This is not directly comparable with the transformer or Qwen cohorts: labels and eligibility differ.",
             size=17, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 10, "Source: paper/main.tex, Table classical; matched 3-fold family-grouped OOF.")

    # 11. DeBERTa result
    slide = new_slide(prs, "Latest DeBERTa study: encouraging ranking, wide uncertainty",
                      "Results · transformer", page=11)
    add_text(slide, 0.95, 1.28, 11.3, 0.38,
             "2,125 clauses  ·  136 positives  ·  9 families  ·  5-fold family-isolated OOF",
             size=16, color=MUTED, bold=True)
    add_bar_chart(slide, 1.15, 2.2, 7.6, [
        ("Equal-weight ECDF ensemble", 0.3711, TEAL),
        ("Partial fine-tuning", 0.3262, BLUE),
        ("Focal loss", 0.3054, ORANGE),
        ("Prevalence baseline", 0.0640, MUTED),
    ], max_value=0.55, baseline=0.064, row_h=0.7, label_w=3.35,
       value_fmt="{:.4f}", highlight="Equal-weight ECDF ensemble")
    add_box(slide, 9.15, 2.0, 3.2, 2.7, fill=PALE_TEAL, outline=PALE_TEAL)
    add_text(slide, 9.45, 2.28, 2.55, 0.55, "26 / 50", size=30, color=TEAL,
             bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 9.45, 2.94, 2.55, 0.85,
             "operational positives among the top 50 ranked clauses",
             size=16, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, 9.36, 4.1, 2.72, 0.42,
             "AP family-cluster interval: 0.315–0.508", size=12,
             color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, 1.1, 5.65, 11.1, 0.62,
             "The interval resamples only nine families; it does not capture training-seed variation or test-set uncertainty.",
             size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 11, "Source: paper/main.tex, Table ranking and fixed-review-budget results.")

    # 12. Hybrid
    slide = new_slide(prs, "The three-model hybrid adds only a small AP increment",
                      "Results · hybrid", page=12)
    add_text(slide, 0.95, 1.25, 11.5, 0.48,
             "Label-compatible subset: 2,097 clauses  ·  116 positives  ·  equal-weight average of within-cohort percentile ranks",
             size=15, color=MUTED, bold=True)
    rows = [
        ("ML logistic regression", 0.4567, BLUE),
        ("V4 semantic ensemble", 0.3373, ORANGE),
        ("DeBERTa-v3-base pilot", 0.2517, TEAL),
        ("V4 + ML", 0.4582, RGBColor(112, 153, 191)),
        ("V4 + DeBERTa + ML", 0.4619, NAVY),
    ]
    add_bar_chart(slide, 0.95, 2.0, 7.9, rows, max_value=0.53, row_h=0.59,
                  label_w=2.9, value_fmt="{:.4f}", highlight="V4 + DeBERTa + ML")
    add_box(slide, 9.1, 2.0, 3.2, 2.78, fill=PALE_ORANGE, outline=PALE_ORANGE)
    add_text(slide, 9.35, 2.3, 2.7, 0.55, "+0.0052 AP", size=27,
             color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 9.4, 2.98, 2.6, 1.1,
             "vs. ML alone on this subset. A modest exploratory difference—not evidence of a reliable gain.",
             size=16, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, 1.02, 5.35, 11.15, 0.78,
             "Caveat: the ML artifact used earlier labels (17 rows missing; 28 conflicts excluded). V4/hybrid design were developed on this cohort.",
             size=15, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 12, "Source: README.md and models/hybrid/README.md. Not independent-test results.")

    # 13. Interpretation
    slide = new_slide(prs, "What can we responsibly conclude from the model studies?",
                      "Synthesis", page=13)
    add_card(slide, 0.88, 1.55, 3.8, 3.55, "1  There is ranking signal",
             "Text and engineered features can prioritize some operationally labeled clauses above prevalence.",
             fill=PALE_BLUE, accent=BLUE, title_size=18, body_size=17)
    add_card(slide, 4.78, 1.55, 3.8, 3.55, "2  Complexity is not a guarantee",
             "The tested embeddings and revision-only features were not consistently stronger; hybrid gains are small.",
             fill=PALE_TEAL, accent=TEAL, title_size=18, body_size=17)
    add_card(slide, 8.68, 1.55, 3.8, 3.55, "3  Cohorts matter",
             "Scores answer different questions when labels, candidate rules, and protocols differ.",
             fill=PALE_ORANGE, accent=ORANGE, title_size=18, body_size=17)
    add_text(slide, 1.0, 5.75, 11.3, 0.62,
             "A useful result is a disciplined shortlist—not a claim that ambiguity has been solved.",
             size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 13, "Source: results/final_summary/conclusions.md; project studies remain cohort-specific.")

    # 14. Limitations
    slide = new_slide(prs, "The main limitations are scientific—not just engineering",
                      "Validity", page=14)
    limitations = [
        ("Construct", "AI-assisted labels; no independent adjudicated expert gold."),
        ("Sampling", "Enriched candidates; scores do not represent every IETF clause."),
        ("Generalization", "Nine families in latest transformer study; one training seed."),
        ("Model selection", "Some comparisons are adaptive development evidence."),
        ("External test", "Reserved TCP and sealed DNS have not been evaluated."),
    ]
    for i, (label, body) in enumerate(limitations):
        y = 1.35 + i * 0.89
        add_box(slide, 1.0, y, 11.25, 0.69, fill=PALE if i % 2 == 0 else WHITE,
                outline=LINE)
        add_text(slide, 1.28, y + 0.17, 2.0, 0.31, label, size=16,
                 color=RED if i in (0, 4) else NAVY, bold=True)
        add_text(slide, 3.38, y + 0.15, 8.5, 0.36, body, size=15, color=INK)
    add_text(slide, 1.15, 6.08, 11.0, 0.42,
             "No claim of validated ambiguity accuracy or deployment readiness is supported yet.",
             size=19, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 14, "Source: paper/main.tex, Threats to Validity.")

    # 15. Use
    slide = new_slide(prs, "The right role is an explainable reviewer aid",
                      "Responsible use", page=15)
    add_text(slide, 1.0, 1.4, 5.0, 0.48, "A human-in-the-loop workflow", size=21,
             color=NAVY, bold=True)
    workflow = [
        ("Rank", "Surface clauses worth a closer read", BLUE),
        ("Inspect", "Return to the exact draft text and context", TEAL),
        ("Decide", "Author or reviewer interprets the requirement", ORANGE),
    ]
    for i, (title, body, accent) in enumerate(workflow):
        y = 2.15 + i * 1.14
        add_circle(slide, 1.1, y, 0.45, accent, str(i + 1), size=16)
        add_text(slide, 1.75, y - 0.01, 1.45, 0.35, title, size=18,
                 color=NAVY, bold=True)
        add_text(slide, 3.15, y - 0.01, 4.0, 0.56, body, size=16, color=INK)
    add_box(slide, 7.8, 1.5, 4.45, 2.55, fill=PALE_RED, outline=PALE_RED)
    add_text(slide, 8.15, 1.78, 3.75, 0.42, "Guardrails", size=21, color=RED, bold=True)
    add_bullet_list(slide, 8.15, 2.3, 3.65, 1.52, [
        "Never auto-reject a draft",
        "Expose uncertainty and provenance",
        "Do not treat generated explanations as proof",
    ], size=13, gap=0.48)
    add_box(slide, 7.8, 4.24, 4.45, 1.32, fill=PALE_TEAL, outline=PALE_TEAL)
    add_text(slide, 8.08, 4.4, 3.9, 0.26, "ILLUSTRATIVE MODEL OUTPUT", size=10,
             color=TEAL, bold=True)
    add_text(slide, 8.08, 4.75, 0.62, 0.34, "YES", size=17, color=TEAL,
             bold=True)
    add_text(slide, 8.78, 4.69, 3.18, 0.66,
             "“When appropriate” leaves the retry condition open.",
             size=13, color=INK)
    add_text(slide, 1.0, 5.9, 11.2, 0.53,
             "A false flag costs reviewer time; a missed flag can create false reassurance.",
             size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 15, "Source: paper/main.tex, Ethical and Practical Considerations.")

    # 16. Next steps
    slide = new_slide(prs, "What evidence would earn a stronger claim?",
                      "Next steps", page=16)
    next_steps = [
        ("01", "Build an expert reference", "Stratified sample; independent double annotation and adjudication.", BLUE),
        ("02", "Freeze the comparison", "Retrain ML on the same frozen labels; repeat finalists across seeds.", TEAL),
        ("03", "Lock the protocol", "Pre-specify model, threshold, code snapshot, and analysis plan.", ORANGE),
        ("04", "Test transfer once", "Evaluate reserved TCP; then open DNS only after the decision rule is fixed.", RED),
    ]
    for i, (number, title, body, accent) in enumerate(next_steps):
        x = 0.9 + i * 3.08
        add_box(slide, x, 1.68, 2.75, 3.47, fill=PALE, outline=LINE)
        add_circle(slide, x + 0.22, 1.94, 0.55, accent, number, size=13)
        add_text(slide, x + 0.22, 2.75, 2.3, 0.75, title, size=17,
                 color=NAVY, bold=True)
        add_text(slide, x + 0.22, 3.68, 2.28, 1.05, body, size=14, color=INK)
    add_text(slide, 1.15, 5.75, 10.95, 0.58,
             "Do not spend the sealed test set to tune the system. Keep it a final transfer check.",
             size=18, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 16, "Source: paper/README.md, author action items; results/final_summary/conclusions.md.")

    # 17. Take-home
    slide = new_slide(prs, "The contribution is a careful path from draft text to review",
                      "Take-home", page=17)
    add_box(slide, 1.05, 1.52, 11.25, 1.2, fill=PALE_BLUE, outline=PALE_BLUE)
    add_text(slide, 1.42, 1.77, 10.5, 0.65,
             "A promising research-stage prioritization aid—not an autonomous ambiguity judge.",
             size=24, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE)
    add_bullet_list(slide, 1.55, 3.22, 10.2, 2.5, [
        "Revision-aware corpus and explicit temporal information boundaries",
        "Family-aware evaluation across classical ML, transformers, and local LLM hybrids",
        "Honest reporting of uncertainty, label limits, and unopened external tests",
    ], size=18, gap=0.83)
    add_text(slide, 1.2, 6.15, 10.9, 0.42,
             "The next milestone is independent evidence—not a more complicated model.",
             size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 17, "Project status and claim boundary: README.md; paper/main.tex.")

    # 18. Discussion and references
    slide = new_slide(prs, "Discussion", "Questions", page=18)
    add_text(slide, 0.95, 1.38, 6.3, 1.08, "Questions?", size=42, color=NAVY, bold=True)
    add_text(slide, 1.0, 2.65, 5.6, 0.82,
             "Where should the boundary sit between automated triage and human judgment?",
             size=22, color=INK)
    add_box(slide, 7.05, 1.43, 5.25, 3.92, fill=PALE, outline=LINE)
    add_text(slide, 7.38, 1.78, 4.6, 0.42, "Selected context", size=19,
             color=TEAL, bold=True)
    refs = [
        "RFC 2119 / RFC 8174 — normative keywords",
        "Davis & Goadrich (2006) — PR curves and ROC",
        "Saito & Rehmsmeier (2015) — PR under imbalance",
        "Pawagi et al. (2025) — RFCScope",
    ]
    add_bullet_list(slide, 7.4, 2.47, 4.5, 2.5, refs, size=14, gap=0.58)
    add_text(slide, 1.0, 6.22, 11.2, 0.38,
             "Thank you  ·  School of Computing, Queen's University", size=17,
             color=MUTED, bold=True)
    add_footer(slide, 18, "Full bibliography and methods: paper/main.tex and paper/references.bib.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return len(prs.slides)


if __name__ == "__main__":
    count = build_deck()
    print(f"Created {OUTPUT} ({count} editable slides)")
